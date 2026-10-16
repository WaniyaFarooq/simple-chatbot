"""
Retrieval-augmented generation helpers.

Pipeline
--------
ingest:    extract (per page) -> clean -> chunk -> embed -> Chroma (one collection per chat)
retrieve:  similarity search -> score floor -> de-duplicate -> light keyword re-rank -> top-k
overview:  for "summarise this document" style questions, sample chunks evenly across the file
"""

import logging
import os
import re
import uuid
from functools import lru_cache
from typing import Dict, List, Optional, Tuple

from docx import Document as DocxDocument
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from . import config

log = logging.getLogger(__name__)

Page = Tuple[Optional[int], str]  # (page number or None, text)
Hit = Dict[str, object]


# ============================================
# EMBEDDINGS + VECTOR STORE
# ============================================


@lru_cache(maxsize=1)
def get_embeddings() -> HuggingFaceEmbeddings:
    """Loaded once, lazily. Normalised so Chroma's cosine scores are meaningful."""
    return HuggingFaceEmbeddings(
        model_name=config.EMBEDDING_MODEL,
        encode_kwargs={"normalize_embeddings": True},
    )


def _collection_name(conversation_id: str) -> str:
    # "_v2": cosine space + normalised embeddings. Older collections are ignored.
    return f"conv_{conversation_id.replace('-', '')}_v2"


def get_vectorstore(conversation_id: str) -> Chroma:
    return Chroma(
        collection_name=_collection_name(conversation_id),
        embedding_function=get_embeddings(),
        persist_directory=config.CHROMA_PATH,
        collection_metadata={"hnsw:space": "cosine"},
    )


# ============================================
# TEXT EXTRACTION
# ============================================


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)   # re-join hyphenated line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pages(file_path: str, extension: str) -> List[Page]:
    extension = extension.lower()

    if extension == ".pdf":
        reader = PdfReader(file_path)
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ValueError("The PDF is password-protected.")
            except ValueError:
                raise
            except Exception:
                raise ValueError("The PDF is password-protected.")

        pages: List[Page] = []
        for number, page in enumerate(reader.pages, start=1):
            text = clean_text(page.extract_text() or "")
            if text:
                pages.append((number, text))
        if not pages:
            raise ValueError(
                "No text found in this PDF. It may be a scanned image (OCR is not supported)."
            )
        return pages

    if extension == ".docx":
        document = DocxDocument(file_path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = clean_text("\n".join(parts))
        if not text:
            raise ValueError("No text found in this document.")
        return [(None, text)]

    if extension == ".txt":
        raw = open(file_path, "rb").read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        text = clean_text(text)
        if not text:
            raise ValueError("The file is empty.")
        return [(None, text)]

    raise ValueError("Unsupported file type. Supported types: PDF, DOCX, TXT")


# ============================================
# INGEST / DELETE
# ============================================


def ingest_document(file_path: str, filename: str, conversation_id: str) -> dict:
    extension = os.path.splitext(filename)[1].lower()
    pages = extract_pages(file_path, extension)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.CHUNK_SIZE,
        chunk_overlap=config.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    batch = uuid.uuid4().hex
    documents: List[Document] = []
    index = 0

    for page, text in pages:
        for chunk in splitter.split_text(text):
            chunk = chunk.strip()
            if len(chunk) < 30:
                continue
            documents.append(
                Document(
                    page_content=chunk,
                    metadata={
                        "conversation_id": conversation_id,
                        "source": filename,
                        "page": page or 0,       # 0 = not applicable (docx / txt)
                        "chunk_index": index,
                        "batch": batch,
                    },
                )
            )
            index += 1

    if not documents:
        raise ValueError("Could not extract any usable text from the document.")

    store = get_vectorstore(conversation_id)
    for start in range(0, len(documents), 64):
        part = documents[start : start + 64]
        store.add_documents(part, ids=[str(uuid.uuid4()) for _ in part])

    # Re-uploading the same filename replaces the previous version.
    try:
        store._collection.delete(
            where={
                "$and": [
                    {"source": {"$eq": filename}},
                    {"batch": {"$ne": batch}},
                ]
            }
        )
    except Exception:
        log.warning("Could not remove previous version of %s", filename, exc_info=True)

    return {"filename": filename, "chunks": len(documents), "pages": len(pages)}


def delete_document(conversation_id: str, filename: str) -> None:
    try:
        get_vectorstore(conversation_id)._collection.delete(
            where={"source": {"$eq": filename}}
        )
    except Exception:
        log.warning("Could not delete vectors for %s", filename, exc_info=True)


def delete_conversation_vectors(conversation_id: str) -> None:
    try:
        get_vectorstore(conversation_id).delete_collection()
    except Exception:
        log.warning("Could not delete collection for %s", conversation_id, exc_info=True)


# ============================================
# RETRIEVAL
# ============================================

_WORD = re.compile(r"[a-z0-9]{4,}")
_OVERVIEW = re.compile(
    r"\b(summari[sz]e|summary|overview|tl;?dr|outline|main points|key points|"
    r"what(?:'s| is) (?:this|the) (?:document|file|pdf|paper|doc|text)|"
    r"about (?:this|the) (?:document|file|pdf|paper|doc))\b",
    re.IGNORECASE,
)


def _terms(text: str) -> set:
    return set(_WORD.findall(text.lower()))


def wants_overview(query: str) -> bool:
    return bool(_OVERVIEW.search(query))


def retrieve_documents(
    conversation_id: str,
    query: str,
    k: Optional[int] = None,
    min_score: Optional[float] = None,
) -> List[Hit]:
    k = k or config.RAG_TOP_K
    min_score = config.RAG_MIN_SCORE if min_score is None else min_score

    store = get_vectorstore(conversation_id)
    raw = store.similarity_search_with_relevance_scores(
        query, k=config.RAG_FETCH_K
    )

    query_terms = _terms(query)
    hits: List[Hit] = []
    seen = set()

    for doc, score in raw:
        if score < min_score:
            continue

        fingerprint = re.sub(r"\W+", "", doc.page_content[:200].lower())
        if fingerprint in seen:
            continue
        seen.add(fingerprint)

        overlap = (
            len(query_terms & _terms(doc.page_content)) / len(query_terms)
            if query_terms
            else 0.0
        )

        hits.append(
            {
                "text": doc.page_content,
                "source": doc.metadata.get("source", "document"),
                "page": doc.metadata.get("page") or None,
                "score": round(float(score), 3),
                "rank": float(score) + 0.1 * overlap,   # keyword boost helps names / numbers
            }
        )

    hits.sort(key=lambda h: h["rank"], reverse=True)
    return hits[:k]


def get_overview_chunks(conversation_id: str, limit: int = 6) -> List[Hit]:
    """Evenly sampled chunks across all documents, for summary-style questions."""
    data = get_vectorstore(conversation_id).get(include=["documents", "metadatas"])
    items = list(zip(data.get("documents") or [], data.get("metadatas") or []))
    if not items:
        return []

    items.sort(key=lambda it: (it[1].get("source", ""), it[1].get("chunk_index", 0)))

    if len(items) > limit:
        step = (len(items) - 1) / (limit - 1)
        items = [items[round(i * step)] for i in range(limit)]

    return [
        {
            "text": text,
            "source": meta.get("source", "document"),
            "page": meta.get("page") or None,
            "score": None,
        }
        for text, meta in items
    ]